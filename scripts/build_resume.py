from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REPO_DIR = Path(__file__).resolve().parent.parent
OUT_DIR = REPO_DIR / 'resume'
DOCX_PATH = OUT_DIR / '李鑫-简历-公开版.docx'

NAVY = '17365D'
MID = '465A70'
LIGHT = 'D9E2F3'
GRAY = '666666'
BLACK = '111111'


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_font(run, size=9.2, bold=False, color=BLACK, latin='Aptos', east='Microsoft YaHei'):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_para(p, before=0, after=0, line=1.0, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = keep
    fmt.keep_with_next = keep


def add_run(p, text, size=9.2, bold=False, color=BLACK, east='Microsoft YaHei'):
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color, east=east)
    return r


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    set_para(p, before=5.0, after=3.2, keep=True)
    pPr = p._p.get_or_add_pPr()
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), NAVY)
    border.append(bottom)
    pPr.append(border)
    add_run(p, title, size=12.0, bold=True, color=NAVY)
    return p


def add_title_row(doc, title, role, date):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(15.1)
    table.columns[1].width = Cm(3.0)
    table.rows[0].cells[0].width = Cm(15.1)
    table.rows[0].cells[1].width = Cm(3.0)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, 0, 0, 0, 0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = table.cell(0, 0).paragraphs[0]
    set_para(p, before=0.4, after=1.0, keep=True)
    add_run(p, title, size=9.8, bold=True)
    if role:
        add_run(p, f'  |  {role}', size=8.9, bold=True, color=MID)
    p2 = table.cell(0, 1).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para(p2, before=0.4, after=1.0, keep=True)
    add_run(p2, date, size=8.7, color=GRAY)
    return table


def add_bullet(doc, text, after=2.0):
    p = doc.add_paragraph(style='List Bullet')
    set_para(p, before=0, after=after, line=1.10)
    p.paragraph_format.left_indent = Cm(0.48)
    p.paragraph_format.first_line_indent = Cm(-0.30)
    add_run(p, text, size=9.35)
    return p


doc = Document()
doc.core_properties.title = '李鑫公开简历'
doc.core_properties.subject = 'AI 应用与全栈开发'
doc.core_properties.author = '李鑫'
doc.core_properties.keywords = 'AI, Full Stack, TypeScript, Java, RAG, Electron, React Native'
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.05)
section.bottom_margin = Cm(0.95)
section.left_margin = Cm(1.35)
section.right_margin = Cm(1.35)
section.header_distance = Cm(0.35)
section.footer_distance = Cm(0.35)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(0)

title_style = styles['Title']
title_style.font.name = 'Aptos Display'
title_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)

# Header
header = doc.add_table(rows=1, cols=2)
header.autofit = False
header.columns[0].width = Cm(18.1)
header.columns[1].width = Cm(0.01)
header.rows[0].cells[0].width = Cm(18.1)
header.rows[0].cells[1].width = Cm(0.01)
for cell in header.rows[0].cells:
    set_cell_margins(cell, 0, 0, 0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

left = header.cell(0, 0)
p = left.paragraphs[0]
p.style = title_style
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_para(p, before=0, after=0)
add_run(p, '李鑫', size=22, bold=True)

p = left.add_paragraph()
set_para(p, before=0.5, after=1.0)
add_run(p, '东北大学计算机科学与技术硕士研究生  |  AI 应用与全栈开发', size=10.1, bold=True, color=NAVY)

p = left.add_paragraph()
set_para(p, before=0, after=0)
add_run(p, '20225802@stu.neu.edu.cn  |  github.com/as2132r2  |  as2132r2.github.io', size=8.8, color=GRAY)

add_section_heading(doc, '个人概要')
p = doc.add_paragraph()
set_para(p, before=0, after=0.8, line=1.05)
add_run(p, '聚焦 AI 应用与桌面端产品开发，具有 Electron 交互、会话状态可靠性、本地数据安全与教学技能工程经验。能从用户问题定位根因，通过自动化测试、CI 与真机走查完成交付。', size=9.4)

add_section_heading(doc, '教育经历')
add_title_row(doc, '东北大学', '计算机科学与技术  硕士研究生在读', '2026 - 至今')
add_title_row(doc, '东北大学', '计算机科学与技术  本科', '2022 - 2026')

add_section_heading(doc, '工程实践')
add_title_row(doc, '薄荷 Agent AI 助手生态', 'AI 应用研发', '2026.06 - 至今')
add_bullet(doc, '在 bohe-ai-consultation 与 juxiang-skills 累计合并 75 个个人 PR。修复会话谱系重启丢失与 SSE 重连吞事件问题，增加本地持久化、LRU 上限、坏文件容错和回归测试；另处理并发 429、会话中毒恢复和跨线程交互串台。')
add_bullet(doc, '提交并迭代学生 AI 创作、心理陪伴与大四生涯规划三类专属工作台，完成作品文件编辑与历史、教练分块回复、停靠对话、本地数据存储和确认后写入机制。')
add_bullet(doc, '开发作文批改与组卷助手技能，支持 OCR 原文确认、Word 原生批注、学生/教师版分离、确定性 HTML 渲染与交付预检；扩展 K12 学段学科 Profile，用测试固化题型、版式和评分约束。')

add_title_row(doc, '贵客松广电内容审核平台', '底座、集成与发布负责人', '2026.08')
add_bullet(doc, '负责共享契约、SQLite/Drizzle 持久化、SSE 事件和配置底座，集成“素材准入 - 稿件生成 - 输出预检 - 三审流转 - AI 参与度追溯”链路，并完成签名会话、固定角色权限与真人审核留痕。')
add_bullet(doc, '建立 GitHub Actions 类型/测试/容器冒烟检查，用 Docker 完成本地与容器化演示，并编写 systemd + Nginx 生产部署、健康检查与凭证安全边界。')

add_title_row(doc, '利欧科研文献管理系统', '本科毕业设计  开源项目二次开发', '2025.12 - 2026.08')
add_bullet(doc, '基于 Spring Boot + Vue 3 整理文献处理与 RAG 问答链路：MinIO 分片上传、Apache Tika 流式解析、Embedding 批处理、Elasticsearch BM25/向量混合检索（0.3/0.7）与 WebSocket 流式回答溯源。')
add_bullet(doc, '通过 Spring Security + JWT 和文档公开属性进行权限过滤，使用 MySQL、Redis、Kafka 和 MinIO 支撑异步处理；完成后端打包、前端生产构建与 12 项单元测试验证。')

add_section_heading(doc, '荣誉与证书')
p = doc.add_paragraph()
set_para(p, before=0, after=0, line=1.05)
add_run(p, '全国大学生数学竞赛国家级三等奖、辽宁省二等奖  |  国家级大创项目优秀结题  |  蓝桥杯 C++ 大学 A 组省三等奖  |  CET-4、CET-6', size=9.1)

# Footer line and page number-like identity.
footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para(p, before=0, after=0)
add_run(p, '李鑫  |  AI 应用与全栈开发', size=7.3, color='8A8A8A')

OUT_DIR.mkdir(parents=True, exist_ok=True)
doc.save(DOCX_PATH)
print(DOCX_PATH)
